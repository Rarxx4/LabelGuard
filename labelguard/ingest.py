"""Turns uploaded files into things the pipeline can use.

Artwork  -> page images (what the AI and the reviewer look at)
          + positioned text lines (what we use to verify the AI).
Documents (rulebooks, specs) -> plain text.
"""
from __future__ import annotations

import io
from pathlib import Path

import pymupdf as fitz
from PIL import Image

from . import config
from .ocr import ocr_image
from .schemas import PageImage, TextLine

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
PT_TO_MM = 25.4 / 72
# Cap height / numeral height of typical sans-serif fonts relative to the em size.
CAP_HEIGHT_RATIO = 0.72


# --------------------------------------------------------------------------
# Artwork
# --------------------------------------------------------------------------
def ingest_artwork(data: bytes, filename: str, out_dir: Path, max_pages: int = 4):
    out_dir.mkdir(parents=True, exist_ok=True)
    ext = Path(filename).suffix.lower()
    warnings: list[str] = []
    if ext == ".pdf":
        pages, lines = _artwork_from_pdf(data, out_dir, max_pages, warnings)
    elif ext in IMAGE_EXT:
        pages, lines = _artwork_from_image(data, out_dir)
    else:
        raise ValueError(f"Unsupported artwork type '{ext}'. Use PDF, PNG or JPG.")
    _number_lines(lines)
    return pages, lines, warnings


def _artwork_from_pdf(data, out_dir, max_pages, warnings):
    doc = fitz.open(stream=data, filetype="pdf")
    if doc.page_count > max_pages:
        warnings.append(f"Artwork has {doc.page_count} pages; only the first {max_pages} were checked.")
    pages, lines = [], []
    for i, page in enumerate(doc):
        if i >= max_pages:
            break
        rect = page.rect
        zoom = config.RENDER_LONG_EDGE_PX / max(rect.width, rect.height)
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        png_path = out_dir / f"page_{i}.png"
        pix.save(png_path)
        pages.append(PageImage(index=i, png_path=str(png_path), width_px=pix.width, height_px=pix.height,
                               width_mm=rect.width * PT_TO_MM, height_mm=rect.height * PT_TO_MM))

        page_lines = _pdf_text_lines(page, i)
        if sum(len(l.text) for l in page_lines) >= config.MIN_TEXT_LAYER_CHARS:
            lines.extend(page_lines)
        else:
            # Typical for print-ready packaging: text converted to outlines.
            warnings.append(f"Page {i + 1}: no usable text layer (outlined text?) - used OCR instead.")
            img = Image.open(png_path)
            mm_per_px = (rect.width * PT_TO_MM) / pix.width
            lines.extend(_ocr_lines(img, i, mm_per_px))
    return pages, lines


def _pdf_text_lines(page, index) -> list[TextLine]:
    W, H = page.rect.width, page.rect.height
    out = []
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            spans = [s for s in line["spans"] if s["text"].strip()]
            if not spans:
                continue
            text = "".join(s["text"] for s in line["spans"]).strip()
            size_pt = max(s["size"] for s in spans)
            x0, y0, x1, y1 = line["bbox"]
            out.append(TextLine(
                line_id="", page=index, text=text, source="pdf",
                bbox=(x0 / W, y0 / H, x1 / W, y1 / H),
                font_size_mm=round(size_pt * PT_TO_MM, 2),
                cap_height_mm=round(size_pt * PT_TO_MM * CAP_HEIGHT_RATIO, 2),
            ))
    return out


def _artwork_from_image(data, out_dir):
    img = Image.open(io.BytesIO(data))
    img.load()
    if img.mode not in ("RGB", "L"):
        bg = Image.new("RGB", img.size, "white")
        bg.paste(img, mask=img.split()[-1] if img.mode in ("RGBA", "LA") else None)
        img = bg
    img = img.convert("RGB")
    scale = config.RENDER_LONG_EDGE_PX / max(img.size)
    if scale < 1:
        img = img.resize((round(img.width * scale), round(img.height * scale)), Image.LANCZOS)
    png_path = out_dir / "page_0.png"
    img.save(png_path)
    page = PageImage(index=0, png_path=str(png_path), width_px=img.width, height_px=img.height)
    return [page], _ocr_lines(img, 0, None)


def _ocr_lines(img: Image.Image, page_index: int, mm_per_px: float | None) -> list[TextLine]:
    out = []
    for r in ocr_image(img):
        cap = None
        if mm_per_px:
            # OCR boxes span ascenders..descenders; numerals are ~60% of that.
            cap = round(r["height_px"] * mm_per_px * 0.6, 2)
        out.append(TextLine(line_id="", page=page_index, text=r["text"], bbox=r["bbox"], source="ocr",
                            cap_height_mm=cap, confidence=r["confidence"]))
    return out


def _number_lines(lines: list[TextLine]):
    lines.sort(key=lambda l: (l.page, round(l.bbox[1], 3), l.bbox[0]))
    for n, line in enumerate(lines, 1):
        line.line_id = f"L{n}"


# --------------------------------------------------------------------------
# Supporting documents
# --------------------------------------------------------------------------
def document_to_text(data: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        doc = fitz.open(stream=data, filetype="pdf")
        text = "\n\n".join(p.get_text("text") for p in doc)
        if len(text.strip()) < 20:  # scanned PDF
            text = "\n".join(
                l["text"] for p in doc
                for l in ocr_image(Image.open(io.BytesIO(p.get_pixmap(dpi=200).tobytes("png"))))
            )
        return text
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    if ext in (".xlsx", ".xlsm", ".xls", ".csv"):
        import pandas as pd
        if ext == ".csv":
            sheets = {"Sheet": pd.read_csv(io.BytesIO(data), dtype=str)}
        else:
            sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=str)
        parts = []
        for name, df in sheets.items():
            parts.append(f"## Sheet: {name}")
            parts.append(df.fillna("").to_csv(sep="|", index=False))
        return "\n".join(parts)
    if ext in IMAGE_EXT:
        return "\n".join(l["text"] for l in ocr_image(Image.open(io.BytesIO(data))))
    return data.decode("utf-8", errors="replace")
